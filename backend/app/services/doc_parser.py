"""Extract plain text from uploaded Excel requirement/QBP workbooks.

The configured model has no native document-ingestion support (Bedrock
Converse `document` content blocks are rejected outright), so every
workbook must be reduced to plain text on our side before it reaches the
model.
"""

import io

from openpyxl import load_workbook

EXCEL_EXTENSIONS = {".xlsx"}


def is_supported(filename: str) -> bool:
    lower = filename.lower()
    return any(lower.endswith(ext) for ext in EXCEL_EXTENSIONS)


def extract_text(filename: str, data: bytes) -> str | None:
    """Return the workbook's content as plain text, or None if no sheet
    contained any data."""
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sections = []
        for sheet in workbook.worksheets:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell).strip() if cell is not None else "" for cell in row]
                if any(cells):
                    lines.append(" | ".join(cells))
            if lines:
                sections.append(f"--- Sheet: {sheet.title} ---\n" + "\n".join(lines))
        return "\n\n".join(sections) if sections else None
    finally:
        workbook.close()
