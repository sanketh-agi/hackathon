"""Extract rule tables from an Excel workbook into JSON.

For every worksheet, locate a header row containing the required columns
`Rule Name` and `Discreption` (and the optional `Workflow`). When both required
columns are present, every data row below the header is emitted as an object;
the worksheet name becomes the JSON key mapping to the list of those objects.
Worksheets missing either required column are skipped entirely.

Usage:
    python tools/extract_rules.py <workbook.xlsx> [-o output.json]

With no -o, the JSON is printed to stdout.
"""

import argparse
import io
import json
import sys

from openpyxl import load_workbook

REQUIRED_COLUMNS = ["Rule Name", "Discreption"]
OPTIONAL_COLUMNS = ["Workflow"]
ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS

# Alternate header spellings/names that map to a canonical column above.
COLUMN_ALIASES = {
    "Rule Name": ["Rule Name", "Attribute Name", "Unique Attribute ID"],
    "Discreption": ["Discreption", "Description", "Attribute Description"],
}

# How many rows from the top to scan when looking for the header row.
HEADER_SCAN_LIMIT = 20


def _norm(value) -> str:
    """Normalise a cell value to a trimmed string for header comparison."""
    if value is None:
        return ""
    return str(value).strip()


def find_header(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    """Scan the top rows for one containing both required column labels.

    Returns (header_row_index, {column_name: column_index}) or None if no row
    within the scan limit contains both required columns.
    """
    for row_idx, row in enumerate(rows[:HEADER_SCAN_LIMIT]):
        labels = {_norm(cell).lower(): col_idx for col_idx, cell in enumerate(row)}

        def resolve(col: str) -> int | None:
            for alias in COLUMN_ALIASES.get(col, [col]):
                if alias.lower() in labels:
                    return labels[alias.lower()]
            return None

        resolved = {col: resolve(col) for col in ALL_COLUMNS}
        if all(resolved[col] is not None for col in REQUIRED_COLUMNS):
            present = {col: idx for col, idx in resolved.items() if idx is not None}
            return row_idx, present
    return None


def extract_sheet(rows: list[tuple]) -> list[dict] | None:
    """Return the list of rule objects for a sheet, or None to skip the sheet."""
    header = find_header(rows)
    if header is None:
        return None
    header_idx, columns = header

    records: list[dict] = []
    for row in rows[header_idx + 1:]:
        rule_name = _norm(row[columns["Rule Name"]]) if columns["Rule Name"] < len(row) else ""
        if not rule_name:
            # Skip blank/separator rows (no rule name on this row).
            continue
        record: dict[str, str] = {}
        for col in ALL_COLUMNS:
            if col not in columns:
                continue  # optional column absent from this sheet
            idx = columns[col]
            record[col] = _norm(row[idx]) if idx < len(row) else ""
        records.append(record)
    return records


def extract_workbook(path: str) -> dict[str, list[dict]]:
    with open(path, "rb") as fh:
        return extract_workbook_bytes(fh.read())


def extract_workbook_bytes(data: bytes) -> dict[str, list[dict]]:
    """Same as extract_workbook, but reads an in-memory .xlsx payload."""
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    result: dict[str, list[dict]] = {}
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            records = extract_sheet(rows)
            if records is not None:
                result[sheet.title] = records
    finally:
        workbook.close()
    return result


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="Extract rule tables from an Excel workbook to JSON.")
    parser.add_argument("workbook", help="Path to the .xlsx workbook")
    parser.add_argument("-o", "--output", help="Write JSON here instead of stdout")
    args = parser.parse_args(argv)

    data = extract_workbook(args.workbook)
    text = json.dumps(data, indent=2, ensure_ascii=False)

    if args.output:
        with open(args.output, "w", encoding="utf-8") as fh:
            fh.write(text)
        print(f"Wrote {len(data)} worksheet(s) to {args.output}")
    else:
        print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
